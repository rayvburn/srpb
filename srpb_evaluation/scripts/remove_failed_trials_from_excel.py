'''
Opens an Excel file with results obtained from SRPB logs, searches for trials that were not
successful (unfinished), and clears the metrics related to those trials so the failed trials are not considered
in filtering.
Note that the Excel spreadsheet must be previously generated with the relevant script,
i.e., 'create_excel_from_results.py', as this script strictly follows the spreadsheet data placement introduced
in the mentioned 'create_excel_from_results.py'

Dependency:
sudo apt install python3-openpyxl
'''

import argparse
import excel_sheet_utils
import shutil

from excel_sheet_utils import get_sheet_val
from excel_sheet_utils import increment_column
from excel_sheet_utils import increment_row
from excel_sheet_utils import is_sheet_cell_empty
from excel_sheet_utils import make_sheet_cell

from openpyxl import load_workbook
from openpyxl.worksheet import Worksheet

from pathlib import Path


def remove_unfinished_trials(
    ws: Worksheet,
    metric_cond_name: str,
    metric_cond_remove_val: float,
    row_init: int,
    col_init: str
):
    # iterate over column with subsequent trials
    col_it = col_init
    # find the row of the target metric
    row_it = row_init
    while not get_sheet_val(ws, row_it, col_it) == metric_cond_name:
        row_it = increment_row(row_it)
    # save
    row_target_metric = row_it

    # iterate along the same row
    while not is_sheet_cell_empty(ws, row_init, col_it):
        remove_trial_if_unsuccessful(ws, row_target_metric, col_it, metric_cond_remove_val)
        col_it = increment_column(col_it)


def remove_trial_if_unsuccessful(
    ws: Worksheet,
    metric_cond_row: int,
    metric_cond_col: str,
    metric_cond_value_to_remove
):
    """
    Assumes that the target metric is located above the metrics to remove
    """
    cond_value = get_sheet_val(ws, metric_cond_row, metric_cond_col)
    if not cond_value == metric_cond_value_to_remove:
        return

    # do not remove the conditional value
    row_it = increment_row(metric_cond_row)
    # iterate over rows to remove all metrics in a given column (related to a given trial)
    while not is_sheet_cell_empty(ws, row_it, metric_cond_col):
        cell = make_sheet_cell(row_it, metric_cond_col)
        ws[cell] = None
        row_it = increment_row(row_it)


if __name__ == "__main__":
    # Ref: https://stackoverflow.com/a/32763023
    cli = argparse.ArgumentParser()
    # positional arguments
    cli.add_argument("input", help="Path to an input spreadsheet with SRPB results")
    cli.add_argument("output", nargs='?', help="Path to an output spreadsheet with SRPB results; empty overwrites the input file")

    # parse the command line
    args = cli.parse_args()

    # location of the input/output files
    input_path = Path(args.input)
    # if optional argument is empty, input file will be overwritten
    if args.output == None:
        output_path = input_path
        # create a temporary backup just in case
        shutil.copyfile(input_path, "/tmp/" + str(input_path.stem) + str(input_path.suffix))
    else:
        output_path = Path(args.output)

    TARGET_METRIC_NAME = 'm_goal'
    TARGET_METRIC_VALUE_REMOVE = 0

    # Load the spreadsheet contents and operate directly on them
    # NOTE: False keeps the functions etc. instead of raw data (True)
    wb = load_workbook(filename=input_path, data_only=False)
    remove_unfinished_trials(
        wb[excel_sheet_utils.SHEET_NAME],
        TARGET_METRIC_NAME,
        TARGET_METRIC_VALUE_REMOVE,
        excel_sheet_utils.RESULT_RAW_INIT_ROW,
        excel_sheet_utils.RESULT_RAW_INIT_COL
    )

    # Save the file
    wb.save(output_path)

    # Copied from 'create_excel_from_results.py'
    print(f'Results saved in: {output_path}')
    print("")
    # When the sheet with the results is not opened and saved by the Excel or LibreOffice Calc, then reading a non-empty
    # cell will probably return None
    # Ref1: https://itecnote.com/tecnote/python-openpyxl-data_onlytrue-returning-none/
    # Ref2: https://groups.google.com/g/openpyxl-users/c/GbBOnOa8g7Y
    print(f'Consider opening the results file and saving it with Excel/LibreOffice Calc (without any modifications).')
    print(f'It will produce cached values based on formulas written (`openpyxl` library is not able to do so).')
    print(f'This is a necessary step when one wants to use the script that creates a LaTeX table from a spreadsheet')
